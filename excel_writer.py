"""
excel_writer.py — writes results.xlsx with all sheets and charts
"""
from __future__ import annotations

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.chart.data_source import AxDataSource, StrRef, StrData, StrVal
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import SeriesLabel
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── Styles ───────────────────────────────────────────────────────────────────────────

_H1   = PatternFill("solid", fgColor="1F3864")
_H2   = PatternFill("solid", fgColor="2E75B6")
_ALT  = PatternFill("solid", fgColor="EBF3FB")
_WARN = PatternFill("solid", fgColor="FFE699")
_ERR  = PatternFill("solid", fgColor="FFC7CE")
_OK   = PatternFill("solid", fgColor="C6EFCE")
_YEL  = PatternFill("solid", fgColor="FFFFC0")

_FW  = Font(bold=True, color="FFFFFF", size=10)
_FD  = Font(bold=True, color="1F3864", size=10)
_FB  = Font(size=10)
_FW2 = Font(color="7F6000", size=10)
_FE  = Font(color="9C0006", size=10)
_FOK = Font(color="276221", size=10)

_s = Side(style="thin", color="B8CCE4")
_B = Border(left=_s, right=_s, top=_s, bottom=_s)

FMT_PCT   = "0.00%"
FMT_MONEY = "#,##0.00"
FMT_INT   = "#,##0"
FMT_DATE  = "DD.MM.YYYY"


def _h(ws, row, n, fill=None, font=None):
    fill = fill or _H1
    font = font or _FW
    for c in range(1, n + 1):
        cell = ws.cell(row, c)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _B
    ws.row_dimensions[row].height = 32


def _cell(ws, r, c, val, fmt=None, fill=None, font=None, align="left"):
    safe = val if not (isinstance(val, float) and np.isnan(val)) else None
    cell = ws.cell(r, c, safe)
    cell.border = _B
    cell.font   = font or _FB
    cell.alignment = Alignment(horizontal=align, vertical="center")
    if fill:
        cell.fill = fill
    if fmt:
        cell.number_format = fmt
    return cell


def _col_width(ws, widths: list[int] | None = None):
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    else:
        for col in ws.columns:
            letter = get_column_letter(col[0].column)
            mx = max((len(str(c.value or "")) for c in col), default=4)
            ws.column_dimensions[letter].width = min(mx + 4, 40)


def _freeze(ws, cell="A2"):
    ws.freeze_panes = cell
    ws.sheet_view.showGridLines = False


# ── Main function ──────────────────────────────────────────────────────────────────

def save_results(path: str, results: dict[str, pd.DataFrame]) -> None:
    """Write results.xlsx: data sheets + charts."""
    wb = Workbook()
    wb.remove(wb.active)

    _write_product_summary(wb, results["product"])
    _write_component_summary(wb, results["component"])
    _write_weekly_detail(wb, results["detail"])
    _write_charts(wb, results)

    wb.save(path)


# ── Data sheets ───────────────────────────────────────────────────────────────────

def _write_product_summary(wb: Workbook, df: pd.DataFrame):
    ws = wb.create_sheet("ProductSummary")
    _freeze(ws)

    labels = ["Product", "Week", "Composite\nYield", "Scrap\nQty",
              "Completed\nDetectors", "Total Scrap\nCost (€)",
              "Scrap /\nDetector (€)", "Status"]
    for ci, lbl in enumerate(labels, 1):
        ws.cell(1, ci, lbl)
    _h(ws, 1, len(labels))

    df_s = df.sort_values(["Product", "Week"])
    for ri, row in enumerate(df_s.itertuples(index=False), 2):
        bg = _ALT if ri % 2 == 0 else None
        _cell(ws, ri, 1, row.Product, align="center", fill=bg)
        _cell(ws, ri, 2, row.Week,    align="center", fill=bg)

        yld = getattr(row, "composite_yield", None)
        yld_fill = (_WARN if pd.notna(yld) and yld < 0.95 else bg)
        _cell(ws, ri, 3, yld if pd.notna(yld) else None, FMT_PCT, fill=yld_fill, align="right")

        bad = getattr(row, "total_bad_qty", None)
        _cell(ws, ri, 4, int(bad) if pd.notna(bad) else None, FMT_INT, fill=bg, align="right")

        det = row.completed_detectors
        has_det = pd.notna(det) and det > 0
        _cell(ws, ri, 5, det if has_det else None, FMT_INT, fill=bg, align="right")

        sc = row.total_scrap_cost
        has_sc = pd.notna(sc)
        _cell(ws, ri, 6, sc if has_sc else None, FMT_MONEY, fill=bg, align="right")

        spd = row.scrap_per_detector
        _cell(ws, ri, 7, spd if pd.notna(spd) else None, FMT_MONEY, fill=bg, align="right")

        if not has_sc:
            status, sf, ff = "No price ⚠", _WARN, _FW2
        elif not has_det:
            status, sf, ff = "No data",     _YEL,  _FB
        else:
            status, sf, ff = "OK ✓",        _OK,   _FOK
        c = ws.cell(ri, 8, status)
        c.fill, c.font, c.border = sf, ff, _B
        c.alignment = Alignment(horizontal="center", vertical="center")

    last_data_row = 1 + len(df_s)
    ws.auto_filter.ref = f"A1:H{max(last_data_row, 1)}"
    summary_start = last_data_row + 2

    title_cell = ws.cell(summary_start, 1, "CUMULATIVE SUMMARY (all weeks)")
    title_cell.fill = _H1
    title_cell.font = _FW
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    title_cell.border = _B
    ws.merge_cells(start_row=summary_start, start_column=1,
                   end_row=summary_start, end_column=8)
    ws.row_dimensions[summary_start].height = 22

    sh = summary_start + 1
    sub_labels = ["Product", "Avg Composite\nYield", "Total\nScrap Qty",
                  "Total Scrap\nCost (€)", "Total\nDetectors",
                  "Scrap /\nDetector (€)", "", ""]
    for ci, lbl in enumerate(sub_labels, 1):
        ws.cell(sh, ci, lbl)
    _h(ws, sh, 6, fill=_H2)
    ws.row_dimensions[sh].height = 32

    products = sorted(df_s["Product"].dropna().unique())
    for pi, product in enumerate(products):
        prow = df_s[df_s["Product"] == product]
        dr = sh + 1 + pi
        bg = _ALT if pi % 2 == 0 else None

        _cell(ws, dr, 1, product, align="center", fill=bg)

        yld_vals = prow["composite_yield"].dropna() if "composite_yield" in prow.columns else pd.Series([], dtype=float)
        avg_yld = yld_vals.mean() if len(yld_vals) else None
        yld_fill = (_WARN if avg_yld is not None and avg_yld < 0.95 else bg)
        _cell(ws, dr, 2, avg_yld, FMT_PCT, fill=yld_fill, align="right")

        bad_vals = prow["total_bad_qty"].dropna() if "total_bad_qty" in prow.columns else pd.Series([], dtype=float)
        total_bad = int(bad_vals.sum()) if len(bad_vals) else None
        _cell(ws, dr, 3, total_bad, FMT_INT, fill=bg, align="right")

        total_cost = prow["total_scrap_cost"].sum() if "total_scrap_cost" in prow.columns else None
        _cell(ws, dr, 4, total_cost if pd.notna(total_cost) else None, FMT_MONEY, fill=bg, align="right")

        det_vals = prow["completed_detectors"].dropna()
        total_det = int(det_vals.sum()) if len(det_vals) else None
        _cell(ws, dr, 5, total_det, FMT_INT, fill=bg, align="right")

        if total_det and pd.notna(total_cost) and total_det > 0:
            _cell(ws, dr, 6, total_cost / total_det, FMT_MONEY, fill=bg, align="right")
        else:
            _cell(ws, dr, 6, None, FMT_MONEY, fill=bg, align="right")

    _col_width(ws, [12, 12, 14, 12, 16, 16, 16, 14])


def _write_component_summary(wb: Workbook, df: pd.DataFrame):
    ws = wb.create_sheet("ComponentSummary")
    _freeze(ws)

    labels = ["Product", "Week", "Item", "Name", "Composite\nYield", "Total Scrap\nCost (€)"]
    for ci, lbl in enumerate(labels, 1):
        ws.cell(1, ci, lbl)
    _h(ws, 1, len(labels), fill=_H2)

    df_s = df.sort_values(["Product", "Week", "Item"])
    for ri, row in enumerate(df_s.itertuples(index=False), 2):
        bg = _ALT if ri % 2 == 0 else None
        _cell(ws, ri, 1, row.Product, align="center", fill=bg)
        _cell(ws, ri, 2, row.Week,    align="center", fill=bg)
        _cell(ws, ri, 3, row.Item,    fill=bg)
        name = getattr(row, "Name", "") or ""
        _cell(ws, ri, 4, name,        fill=bg)

        yld = row.composite_yield
        yld_fill = (_WARN if pd.notna(yld) and yld < 0.95 else bg)
        _cell(ws, ri, 5, yld if pd.notna(yld) else None, FMT_PCT, fill=yld_fill, align="right")
        _cell(ws, ri, 6, row.total_scrap_cost, FMT_MONEY, fill=bg, align="right")

    ws.auto_filter.ref = f"A1:F{max(1 + len(df_s), 1)}"
    _col_width(ws, [12, 12, 12, 22, 14, 16])


def _write_weekly_detail(wb: Workbook, df: pd.DataFrame):
    ws = wb.create_sheet("WeeklyDetail")
    _freeze(ws)

    labels = ["Product", "Week", "Item", "Name", "Op No", "Task Description",
              "Good", "Bad", "Op Yield", "Accum.\nCost (€)", "Scrap\nCost (€)",
              "Rows", "Date From", "Date To"]
    cols   = ["Product", "Week", "Item", "Name", "OperationNo", "TaskDescription",
              "weekly_good", "weekly_bad", "operation_yield", "AccumulatedCost",
              "operation_scrap_cost", "row_count", "date_min", "date_max"]
    for ci, lbl in enumerate(labels, 1):
        ws.cell(1, ci, lbl)
    _h(ws, 1, len(labels))

    df_s = df.sort_values(["Product", "Week", "Item", "OperationNo"])
    available = [c for c in cols if c in df_s.columns]

    for ri, row in enumerate(df_s[available].itertuples(index=False), 2):
        bg = _ALT if ri % 2 == 0 else None
        for ci, (col, val) in enumerate(zip(available, row), 1):
            fmt, align, fill, font = None, "left", bg, None
            if col in ("Product", "Week"):
                align = "center"
            elif col in ("OperationNo", "weekly_good", "weekly_bad", "row_count"):
                fmt, align = FMT_INT, "right"
            elif col == "operation_yield":
                fmt, align = FMT_PCT, "right"
                if pd.notna(val) and val < 0.95:
                    fill = _WARN
            elif col in ("AccumulatedCost", "operation_scrap_cost"):
                fmt, align = FMT_MONEY, "right"
                if pd.isna(val):
                    fill, font, val = _WARN, _FW2, "⚠ no price"
            elif col in ("date_min", "date_max"):
                fmt = FMT_DATE
            safe = val if not (isinstance(val, float) and np.isnan(val)) else None
            if isinstance(val, str):
                safe = val
            _cell(ws, ri, ci, safe, fmt=fmt, fill=fill, font=font, align=align)

    ws.auto_filter.ref = f"A1:N{max(1 + len(df_s), 1)}"
    _col_width(ws, [10, 10, 12, 22, 7, 28, 8, 8, 9, 12, 12, 6, 12, 12])


# ── Charts ────────────────────────────────────────────────────────────────────────

def _write_charts(wb: Workbook, results: dict[str, pd.DataFrame]):
    ws = wb.create_sheet("Charts")
    ws.sheet_view.showGridLines = False

    prod = results["product"].copy()
    comp = results["component"].copy()

    if prod.empty:
        ws["A1"] = "No data for charts"
        return

    weeks    = sorted(prod["Week"].dropna().unique())
    products = sorted(prod["Product"].dropna().unique())
    items    = sorted(comp["Item"].dropna().unique())

    row = 1
    ws.cell(row, 1, "Total Scrap Cost (€)")
    ws.cell(row, 1).fill = _H2
    ws.cell(row, 1).font = _FW
    ws.row_dimensions[row].height = 20
    row += 1

    ws.cell(row, 1, "Week")
    ws.cell(row, 1).font = _FD
    for ci, p in enumerate(products, 2):
        ws.cell(row, ci, p).font = _FD
    _h(ws, row, 1 + len(products), fill=_H2)
    scrap_header_row = row
    row += 1

    scrap_data_start = row
    for w in weeks:
        ws.cell(row, 1, w.split("-", 1)[1] if "-W" in w else w)
        for ci, p in enumerate(products, 2):
            mask = (prod["Product"] == p) & (prod["Week"] == w)
            val = prod.loc[mask, "total_scrap_cost"].values
            ws.cell(row, ci, float(val[0]) if len(val) and pd.notna(val[0]) else None)
            ws.cell(row, ci).number_format = FMT_MONEY
        row += 1
    scrap_data_end = row - 1

    row += 1
    ws.cell(row, 1, "Scrap per Detector (€)")
    ws.cell(row, 1).fill = _H2
    ws.cell(row, 1).font = _FW
    ws.row_dimensions[row].height = 20
    row += 1

    ws.cell(row, 1, "Week")
    for ci, p in enumerate(products, 2):
        ws.cell(row, ci, p)
    _h(ws, row, 1 + len(products), fill=_H2)
    spd_header_row = row
    row += 1

    spd_data_start = row
    for w in weeks:
        ws.cell(row, 1, w.split("-", 1)[1] if "-W" in w else w)
        for ci, p in enumerate(products, 2):
            mask = (prod["Product"] == p) & (prod["Week"] == w)
            val = prod.loc[mask, "scrap_per_detector"].values
            ws.cell(row, ci, float(val[0]) if len(val) and pd.notna(val[0]) else None)
            ws.cell(row, ci).number_format = FMT_MONEY
        row += 1
    spd_data_end = row - 1

    row += 1
    ws.cell(row, 1, "Composite Yield by Component")
    ws.cell(row, 1).fill = _H2
    ws.cell(row, 1).font = _FW
    ws.row_dimensions[row].height = 20
    row += 1

    ws.cell(row, 1, "Week")
    for ci, it in enumerate(items, 2):
        ws.cell(row, ci, it)
    _h(ws, row, 1 + len(items), fill=_H2)
    yld_header_row = row
    row += 1

    yld_data_start = row
    for w in weeks:
        ws.cell(row, 1, w.split("-", 1)[1] if "-W" in w else w)
        for ci, it in enumerate(items, 2):
            mask = (comp["Item"] == it) & (comp["Week"] == w)
            val = comp.loc[mask, "composite_yield"].values
            if len(val) and pd.notna(val[0]):
                c = ws.cell(row, ci, float(val[0]))
                c.number_format = FMT_PCT
            else:
                ws.cell(row, ci, None)
        row += 1
    yld_data_end = row - 1

    _col_width(ws, [12] + [16] * max(len(products), len(items), 1))

    if len(weeks) < 2:
        ws.cell(row + 1, 1, "At least 2 weeks of data required for charts")
        return

    def _str_categories(chart, col: int, start_row: int, end_row: int):
        col_letter = get_column_letter(col)
        ref = f"'Charts'!${col_letter}${start_row}:${col_letter}${end_row}"
        values = [str(ws.cell(r, col).value) for r in range(start_row, end_row + 1)]
        cache = StrData(
            ptCount=len(values),
            pt=[StrVal(idx=i, v=v) for i, v in enumerate(values)],
        )
        for ser in chart.series:
            ser.cat = AxDataSource(strRef=StrRef(f=ref, strCache=cache))
        chart.x_axis.delete = False
        chart.y_axis.delete = False

    def _data_labels(num_fmt: str | None = None,
                     pos: str | None = None) -> DataLabelList:
        lbl = DataLabelList()
        lbl.showVal      = True
        lbl.showLegendKey = False
        lbl.showCatName  = False
        lbl.showSerName  = False
        if num_fmt:
            lbl.numFmt = num_fmt
        if pos:
            lbl.dLblPos = pos
        return lbl

    CHART_W, CHART_H, STEP = 22, 10, 22
    chart_row = scrap_header_row

    chart1 = BarChart()
    chart1.type               = "col"
    chart1.grouping           = "clustered"
    chart1.title              = "Total Scrap Cost by Week"
    chart1.style              = 10
    chart1.y_axis.title       = "€"
    chart1.y_axis.numFmt      = "#,##0.00"
    chart1.y_axis.scaling.min = 0
    chart1.x_axis.title       = "Week"
    chart1.height             = CHART_H
    chart1.width              = CHART_W
    chart1.dLbls              = _data_labels("#,##0.00")

    for ci, p in enumerate(products, 2):
        data = Reference(ws, min_col=ci, max_col=ci,
                         min_row=scrap_header_row, max_row=scrap_data_end)
        chart1.add_data(data, titles_from_data=True)

    _str_categories(chart1, 1, scrap_data_start, scrap_data_end)
    ws.add_chart(chart1, f"H{chart_row}")
    chart_row += STEP

    chart2 = BarChart()
    chart2.type               = "col"
    chart2.grouping           = "clustered"
    chart2.title              = "Scrap per Detector by Week"
    chart2.style              = 10
    chart2.y_axis.title       = "€ / unit"
    chart2.y_axis.numFmt      = "#,##0.00"
    chart2.y_axis.scaling.min = 0
    chart2.x_axis.title       = "Week"
    chart2.height             = CHART_H
    chart2.width              = CHART_W
    chart2.dLbls              = _data_labels("#,##0.00")

    for ci, p in enumerate(products, 2):
        data = Reference(ws, min_col=ci, max_col=ci,
                         min_row=spd_header_row, max_row=spd_data_end)
        chart2.add_data(data, titles_from_data=True)

    _str_categories(chart2, 1, spd_data_start, spd_data_end)
    ws.add_chart(chart2, f"H{chart_row}")
    chart_row += STEP

    item_product = dict(zip(comp["Item"].astype(str), comp["Product"].astype(str)))
    item_col     = {it: ci for ci, it in enumerate(items, 2)}

    prod_items: dict[str, list[str]] = {}
    for it in items:
        prod_items.setdefault(item_product.get(it, "?"), []).append(it)

    for p in sorted(prod_items):
        chart = LineChart()
        chart.title              = f"Composite Yield — {p}"
        chart.style              = 10
        chart.y_axis.title       = "Yield"
        chart.y_axis.numFmt      = "0.0%"
        chart.y_axis.scaling.min = 0.7
        chart.y_axis.scaling.max = 1.05
        chart.y_axis.majorUnit   = 0.05
        chart.x_axis.title       = "Week"
        chart.height             = CHART_H
        chart.width              = CHART_W

        for it in prod_items[p]:
            data = Reference(ws, min_col=item_col[it], max_col=item_col[it],
                             min_row=yld_header_row, max_row=yld_data_end)
            chart.add_data(data, titles_from_data=True)

        _str_categories(chart, 1, yld_data_start, yld_data_end)
        for ser in chart.series:
            ser.smooth = False
        ws.add_chart(chart, f"H{chart_row}")
        chart_row += STEP
